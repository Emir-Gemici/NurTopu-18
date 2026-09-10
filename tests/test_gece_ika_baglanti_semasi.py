"""NÖBETÇİ — İKA renk kodlu bağlantı şeması (Y-3.3.1-M01 İKA kolu; FDR sprinti H3 · T217).

Ön kayıt (results/fdr_sprint_20260910/ON_KAYIT.md H3; D-439 §3): İKA aviyoniği
TEDARİK EDİLMEDİ, ölçülmüş İKA topolojisi YOKTUR. Şema DDR TASARIMINDAN çizilir:
"her hat 'DDR tasarımı, tedarik yok' etiketli; kaynak DDR BOM/mimari; ölçülmüş hat 0".
İHA şemasının kalıbı (tests/test_gece_baglanti_semasi.py) aynen izlenir; farkı, bu
şemada DÜZ ÇİZGİNİN YASAK olmasıdır — ölçülmeyen hat ölçülmüş gibi çizilemez.

  T1  SVG geçerli XML, kök <svg>, genişlik 1100, <title> var
  T2  legend'de zorunlu türler (ana güç, regüle güç, GND, motor, UART, I2C, CSI,
      enkoder, GPIO, analog) ayrı renkle var
  T3  çizilen her hat legend'deki bir rengi kullanır ve data-tur rengi legend ile eş
  T4  kaynaklar.md'de her hat satırının kaynak hücresi bir DDR şekil/tablo atfı taşır
      ("Şekil 3.x.y.z" / "Tablo 3.x.y.z"); kaynaksız hat çizilmez
  T5  SVG'deki hat kümesi == md'deki hat kümesi == üreticinin HATLAR kümesi;
      md "Sayım" satırı aynı sayıyı söyler
  T6  ölçülmüş hat 0: her hat KESİK; her durum hücresi "DDR tasarımı, tedarik yok" ile
      başlar; düz çizgi ancak durum "ÖLÇÜLDÜ" ile başlıyorsa mümkün (bugün hiç yok)
  T7  atıf yapılan her D-kaydı REVIEW_LOG.md'de VAR; DDR şekil/tablo atıfları
      kullanıcının DDR docx'lerinde başlık olarak VAR (docx yoksa/python-docx yoksa
      bu alt kapı ATLANIR — ÖLÇÜLEMEDİ ≠ GEÇTİ, atlanma görünür)
  T8  isim yok: commit yazarının adı ve yönergenin yasakladığı adlar geçmez
  T9  drift kilidi: diskteki iki dosya üreticinin çıktısıyla bayt-eş
  T9b FDR BOLUM_3 §3.3.1.1 şemayı gömüyor, eski "şema yoktur" cümlesi üstü çizili,
      şema satırı "DDR tasarımı" ve "ölçülmüş" sayısını söylüyor; YONERGE_HARITASI
      Y-3.3.1-M01 satırı ika_baglanti.svg'yi anıyor
  T10 kural fonksiyonları sentetik girdide de yakalar (mutasyon: legend dışı renk,
      kaynaksız satır, ölçülmemiş hat düz çizilmiş, DDR etiketsiz durum)

Bu dosya SVG'yi güzel bulmaz; yalnız sayılabilir olanı sayar.
"""

from __future__ import annotations

import importlib.util as _u
import re
import subprocess
import unittest
import xml.etree.ElementTree as ET
from pathlib import Path

KOK = Path(__file__).resolve().parents[1]
SEKILLER = KOK / "FDR" / "sekiller"
SVG = SEKILLER / "ika_baglanti.svg"
MD = SEKILLER / "ika_baglanti_kaynaklar.md"
BETIK = KOK / "scripts" / "fdr_ika_baglanti_svg.py"
REVIEW_LOG = KOK / "REVIEW_LOG.md"
BOLUM_3 = KOK / "FDR" / "BOLUM_3_SISTEM_SPESIFIKASYONLARI.md"
YONERGE = KOK / "FDR" / "YONERGE_HARITASI.md"

# Kullanıcının DDR belgeleri: repoda TUTULMAZ (AGENTS.md §2 salt okuma). Ana ağaçta
# `rapor/` altındadır; worktree'de yoktur. Bulunamazsa T7b atlanır (görünür).
DDR_DOCX = {
    "avi": "3.3.1-2-3-4_SON.docx",
    "yaz": "3.3.5_Yazilim_Mimarisi_MERGED_duzeltmeler.docx",
    "mek": "NAVITAS_IKA_DDR.docx",
}

SVG_NS = "{http://www.w3.org/2000/svg}"
ZORUNLU_TURLER = {"guc_hv", "guc", "guc_gnd", "motor", "uart", "i2c", "csi", "enk", "gpio", "analog"}
DDR_ETIKETI = "DDR tasarımı, tedarik yok"
DDR_ATIF = re.compile(r"(Şekil|Tablo)\s+3\.\d+\.\d+\.\d+")
D_KAYIT = re.compile(r"\bD-(\d{3})\b")
HAT_SATIRI = re.compile(r"^\|\s*(K\d{2})\s*\|")
SAYIM = re.compile(r"Hat:\s*\*\*(\d+)\*\*\s*·\s*ölçülmüş \(düz\):\s*\*\*(\d+)\*\*")
YASAK_ISIM = re.compile(r"Halil|Yunus|Emir|Doğukan", re.I)


def _oku(p: Path) -> str:
    return p.read_text(encoding="utf-8").replace("\r\n", "\n")


# ── Saf kural fonksiyonları (sentetik girdide de koşar) ──────────────────────
def legend_renkleri(kok: ET.Element) -> dict:
    """legend grubundaki data-hat -> stroke."""
    for g in kok.iter(SVG_NS + "g"):
        if g.get("id") == "legend":
            return {el.get("data-hat"): el.get("stroke")
                    for el in g.iter() if el.get("data-hat") and el.get("stroke")}
    return {}


def hatlar(kok: ET.Element) -> list:
    """class='hat' elemanları: (id, tür, stroke, kesik_mi)."""
    return [(el.get("data-hat"), el.get("data-tur"), el.get("stroke"),
             el.get("stroke-dasharray") is not None)
            for el in kok.iter() if "hat" in (el.get("class") or "").split()]


def legend_disi_renkler(kok: ET.Element) -> list:
    renkler = set(legend_renkleri(kok).values())
    return [(i, r) for i, _, r, _ in hatlar(kok) if r not in renkler]


def md_hat_satirlari(metin: str) -> dict:
    """md tablosundan hat id -> (durum, kaynak)."""
    out = {}
    for satir in metin.splitlines():
        m = HAT_SATIRI.match(satir)
        if not m:
            continue
        hucreler = [h.strip() for h in satir.strip().strip("|").split("|")]
        # | hat | ad | nereden → nereye | tür | durum | kaynak |
        out[m.group(1)] = ("", "") if len(hucreler) < 6 else (hucreler[4], hucreler[5])
    return out


def kaynaksiz_satirlar(metin: str) -> list:
    return [hid for hid, (_, kaynak) in md_hat_satirlari(metin).items() if not DDR_ATIF.search(kaynak)]


def ddr_etiketsiz_satirlar(metin: str) -> list:
    return [hid for hid, (durum, _) in md_hat_satirlari(metin).items()
            if not (durum.startswith(DDR_ETIKETI) or durum.startswith("ÖLÇÜLDÜ"))]


def duz_cizilen_olculmemis(kok: ET.Element, metin: str) -> list:
    """Düz (kesiksiz) çizilen hat yalnız durum 'ÖLÇÜLDÜ' ile başlıyorsa mümkündür."""
    durumlar = {hid: d for hid, (d, _) in md_hat_satirlari(metin).items()}
    return [(hid, durumlar.get(hid, "")[:40]) for hid, _, _, kesik in hatlar(kok)
            if not kesik and not durumlar.get(hid, "").startswith("ÖLÇÜLDÜ")]


def _modul():
    sp = _u.spec_from_file_location("fdr_ika_baglanti_svg", BETIK)
    m = _u.module_from_spec(sp)
    sp.loader.exec_module(m)
    return m


def _ddr_basliklari():
    """DDR docx'lerinden 'Şekil/Tablo 3.x.y.z' başlık numaraları; docx yoksa None."""
    try:
        from docx import Document  # python-docx (venv'de var; Jetson'da olmayabilir)
    except ImportError:
        return None
    adaylar = [KOK / "rapor", KOK.parent.parent.parent / "rapor"]  # ana ağaç ↔ worktree
    basliklar = set()
    bulundu = 0
    for ad in DDR_DOCX.values():
        for kok in adaylar:
            yol = kok / ad
            if yol.is_file():
                d = Document(str(yol))
                metin = "\n".join(p.text for p in d.paragraphs)
                for t in d.tables:
                    for r in t.rows:
                        metin += "\n" + " | ".join(c.text for c in r.cells)
                basliklar |= {m.group(0).replace("  ", " ") for m in DDR_ATIF.finditer(metin)}
                bulundu += 1
                break
    return basliklar if bulundu == len(DDR_DOCX) else None


class Sema(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.svg_metin = _oku(SVG)
        cls.md_metin = _oku(MD)
        cls.kok = ET.fromstring(cls.svg_metin)

    def test_T1_svg_xml_gecerli_1100px(self):
        self.assertEqual(self.kok.tag, SVG_NS + "svg")
        self.assertEqual(self.kok.get("width"), "1100")
        self.assertTrue((self.kok.get("viewBox") or "").startswith("0 0 1100 "))
        self.assertIsNotNone(self.kok.find(SVG_NS + "title"))
        self.assertIn(DDR_ETIKETI, self.kok.find(SVG_NS + "title").text or "")

    def test_T2_legend_zorunlu_turler_ayri_renk(self):
        leg = legend_renkleri(self.kok)
        eksik = ZORUNLU_TURLER - set(leg)
        self.assertFalse(eksik, f"legend'de eksik tür: {sorted(eksik)}")
        zorunlu_renkler = [leg[t] for t in sorted(ZORUNLU_TURLER)]
        self.assertEqual(len(set(zorunlu_renkler)), len(ZORUNLU_TURLER),
                         f"zorunlu türler ayrı renk olmalı: {leg}")

    def test_T3_her_hat_legend_renginde(self):
        self.assertGreater(len(hatlar(self.kok)), 0, "hiç hat çizilmemiş")
        self.assertEqual(legend_disi_renkler(self.kok), [])
        leg = legend_renkleri(self.kok)
        for hid, tur, renk, _ in hatlar(self.kok):
            self.assertEqual(leg.get(tur), renk, f"{hid}: data-tur={tur} rengi legend ile uyuşmuyor")

    def test_T4_md_her_hatta_ddr_atfi(self):
        satirlar = md_hat_satirlari(self.md_metin)
        self.assertGreater(len(satirlar), 0)
        self.assertEqual(kaynaksiz_satirlar(self.md_metin), [])
        self.assertEqual(ddr_etiketsiz_satirlar(self.md_metin), [])

    def test_T5_svg_md_uretici_hat_kumeleri_esit(self):
        svg_ids = {hid for hid, _, _, _ in hatlar(self.kok)}
        md_ids = set(md_hat_satirlari(self.md_metin))
        uretici_ids = {h["id"] for h in _modul().HATLAR}
        self.assertEqual(svg_ids, md_ids,
                         f"yalnız SVG'de: {sorted(svg_ids - md_ids)} · yalnız md'de: {sorted(md_ids - svg_ids)}")
        self.assertEqual(md_ids, uretici_ids)
        self.assertGreaterEqual(len(md_ids), 12, "DDR şeması 12'den az hatla anlatılamaz")
        m = SAYIM.search(self.md_metin)
        self.assertIsNotNone(m, "md 'Sayım' satırı yok")
        self.assertEqual(int(m.group(1)), len(md_ids))

    def test_T6_olculmus_hat_sifir_hepsi_kesik(self):
        self.assertEqual(duz_cizilen_olculmemis(self.kok, self.md_metin), [])
        kesik_olmayan = [hid for hid, _, _, kesik in hatlar(self.kok) if not kesik]
        self.assertEqual(kesik_olmayan, [], "İKA aviyoniği tedarik edilmedi: düz çizgi olamaz")
        self.assertEqual(int(SAYIM.search(self.md_metin).group(2)), 0, "md 'ölçülmüş (düz)' 0 olmalı")
        for hid, (durum, _) in md_hat_satirlari(self.md_metin).items():
            self.assertTrue(durum.startswith(DDR_ETIKETI), f"{hid}: durum '{durum[:40]}'")

    def test_T7_atiflar_gercek(self):
        mevcut = set(re.findall(r"^### D-(\d{3})\b", _oku(REVIEW_LOG), flags=re.M))
        atifli = set(D_KAYIT.findall(self.md_metin)) | set(D_KAYIT.findall(self.svg_metin))
        self.assertTrue(atifli, "hiç D-kaydı atfı yok (D-439 beklenir)")
        yok = sorted(a for a in atifli if a not in mevcut)
        self.assertEqual(yok, [], f"REVIEW_LOG'da olmayan D-kaydına atıf: {yok}")

    def test_T7b_ddr_sekil_tablo_atiflari_docxte_var(self):
        basliklar = _ddr_basliklari()
        if basliklar is None:
            self.skipTest("DDR docx'leri (rapor/) ya da python-docx yok — ÖLÇÜLEMEDİ, GEÇTİ sayılmaz")
        atifli = {m.group(0) for m in DDR_ATIF.finditer(self.md_metin)}
        self.assertTrue(atifli)
        yok = sorted(a for a in atifli if a not in basliklar)
        self.assertEqual(yok, [], f"DDR belgelerinde başlık olarak bulunmayan atıf: {yok}")

    def test_T8_isim_yok(self):
        metin = self.md_metin + "\n" + self.svg_metin
        self.assertIsNone(YASAK_ISIM.search(metin), "şemaya kişi adı sızmış")
        self.assertIsNone(re.search(r"(?i)(operat[öo]r|pilot|kaptan)\s*[:(]\s*[A-ZÇĞİÖŞÜ][a-zçğıöşü]+", metin))
        try:
            q = subprocess.run(["git", "log", "-1", "--format=%an"], cwd=KOK,
                               capture_output=True, text=True, timeout=30)
            yazar = q.stdout.strip()
        except OSError:
            yazar = ""
        if not yazar:
            self.skipTest("git yazarı okunamadı")
        for parca in re.split(r"[\s\-_.]+", yazar):
            if len(parca) >= 3:
                self.assertNotRegex(metin, rf"\b{re.escape(parca)}\b", f"isim geçiyor: {parca}")

    def test_T9_drift_kilidi_uretici_ile_ayni(self):
        svg, md = _modul().uret()
        self.assertEqual(svg, self.svg_metin, "ika_baglanti.svg üreticiden sapmış — betiği koş ve commit'le")
        self.assertEqual(md, self.md_metin, "ika_baglanti_kaynaklar.md üreticiden sapmış")

    def test_T9b_bolum3_ve_yonerge_semayi_taniyor(self):
        b3 = _oku(BOLUM_3)
        self.assertIn("sekiller/ika_baglanti.svg", b3, "BOLUM_3 İKA şemasını gömmüyor")
        self.assertTrue(SVG.is_file())
        kesit = b3[b3.index("sekiller/ika_baglanti.svg"):][:1500]
        self.assertIn("DDR tasarımı", kesit, "şema satırı kesik çizginin anlamını (DDR tasarımı) söylemeli")
        self.assertRegex(kesit, r"ölçülmüş", "şema satırı ölçülmüş hat sayısını söylemeli")
        self.assertIn("ika_baglanti_kaynaklar.md", kesit)
        # Eski cümle silinmedi, üstü çizildi (FDR kuralı: eski metin ~~çizili~~ + tarih)
        self.assertRegex(b3, r"~~\*\*BEYAN GEREKLİ \(insan\):\*\* İKA için bağlantı şeması",
                         "eski 'şema yoktur' cümlesi üstü çizili durmalı")
        self.assertIn("ika_baglanti_kaynaklar.md", self.svg_metin, "SVG kaynak dosyasına atıf yapmalı")
        y = _oku(YONERGE)
        satir = next((s for s in y.splitlines() if s.startswith("| `Y-3.3.1-M01`")), "")
        self.assertIn("ika_baglanti.svg", satir, "YONERGE_HARITASI Y-3.3.1-M01 satırı İKA şemasını anmıyor")


class SentetikMutasyon(unittest.TestCase):
    """Kural fonksiyonları gerçek dosyaya değil kurala bağlıdır (D-370 sınıfı)."""

    SVG = ('<svg xmlns="http://www.w3.org/2000/svg" width="1100" viewBox="0 0 1100 10">'
           '<g id="legend"><line data-hat="guc_hv" stroke="#7f0000"/><line data-hat="uart" stroke="#1565c0"/></g>'
           '<polyline class="hat" data-hat="K01" data-tur="guc_hv" stroke="#7f0000" points="0,0 1,1" stroke-dasharray="8,5"/>'
           '<polyline class="hat" data-hat="K02" data-tur="uart" stroke="#1565c0" points="0,0 1,1" stroke-dasharray="8,5"/>'
           '{ek}</svg>')
    MD = ("| hat | ad | nereden → nereye | tür | durum | kaynak |\n|---|---|---|---|---|---|\n"
          "| K01 | a | b → c | `guc_hv` | DDR tasarımı, tedarik yok — x | DDR aviyonik Şekil 3.3.1.5 |\n"
          "| K02 | a | b → c | `uart` | DDR tasarımı, tedarik yok — y | DDR yazılım Tablo 3.3.5.9 |\n{ek}")

    def test_temiz_gecer(self):
        kok = ET.fromstring(self.SVG.format(ek=""))
        md = self.MD.format(ek="")
        self.assertEqual(legend_disi_renkler(kok), [])
        self.assertEqual(kaynaksiz_satirlar(md), [])
        self.assertEqual(ddr_etiketsiz_satirlar(md), [])
        self.assertEqual(duz_cizilen_olculmemis(kok, md), [])

    def test_legend_disi_renk_yakalanir(self):
        kok = ET.fromstring(self.SVG.format(
            ek='<polyline class="hat" data-hat="K03" data-tur="motor" stroke="#ff00ff" points="0,0 1,1" stroke-dasharray="8,5"/>'))
        self.assertEqual(legend_disi_renkler(kok), [("K03", "#ff00ff")])

    def test_kaynaksiz_satir_yakalanir(self):
        md = self.MD.format(ek="| K03 | a | b → c | `motor` | DDR tasarımı, tedarik yok | ekipten duyuldu |\n")
        self.assertEqual(kaynaksiz_satirlar(md), ["K03"])

    def test_ddr_etiketsiz_durum_yakalanir(self):
        md = self.MD.format(ek="| K03 | a | b → c | `motor` | tasarımdan çizildi | DDR aviyonik Şekil 3.3.1.5 |\n")
        self.assertEqual(ddr_etiketsiz_satirlar(md), ["K03"])

    def test_olculmemis_hat_duz_cizilince_yakalanir(self):
        svg = self.SVG.format(ek="").replace(
            'data-hat="K02" data-tur="uart" stroke="#1565c0" points="0,0 1,1" stroke-dasharray="8,5"',
            'data-hat="K02" data-tur="uart" stroke="#1565c0" points="0,0 1,1"')
        kotu = duz_cizilen_olculmemis(ET.fromstring(svg), self.MD.format(ek=""))
        self.assertEqual([k[0] for k in kotu], ["K02"])
        # ÖLÇÜLDÜ durumu düz çizgiye izin verir (İHA şemasındaki kural; İKA'da bugün yok)
        md_olculdu = self.MD.format(ek="").replace("DDR tasarımı, tedarik yok — y", "ÖLÇÜLDÜ (tezgâh)")
        self.assertEqual(duz_cizilen_olculmemis(ET.fromstring(svg), md_olculdu), [])


if __name__ == "__main__":
    unittest.main()
