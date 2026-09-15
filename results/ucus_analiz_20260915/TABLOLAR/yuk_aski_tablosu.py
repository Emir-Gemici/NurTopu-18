"""Ek sayfa: 4S 13 Ah + 9x5x3 icin eklenen yuk -> aski gazi / akim / guc / ucus suresi.
Kalibrasyon performans_zarfi/tablolar.json'dan (k_T, k_P TOPLAM 4 motor; AUW 4,0 kg kullanici beyani).
Kosum: python yuk_aski_tablosu.py  -> xlsx'e 'Y_Yuk_Aski_Gazi' sayfasi + png/46_Y_Yuk_Aski_Gazi.png (tablolar_uret.py sonrasi kosulmali)."""
import json, io, math, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
KOK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
k = json.load(io.open(os.path.join(KOK, "performans_zarfi", "tablolar.json"), encoding="utf-8"))["kalibrasyon"]["4S_13Ah_9x5x3"]
kT, kP, KV, R, E, G = k["k_T_N_per_rpm2"], k["k_P_W_per_rpm3"], 980.0, 0.0304, 0.65, 9.80665
AUW0, KAP_AH, KULL = 4.0, 13.0, 0.80
def tho(a): return ((2*E*a - (E-1))**2 - (1-E)**2) / (4*E)
def aski(m, V0):
    T = m*G; rpm = math.sqrt(T/kT); P = kP*rpm**3; V = V0
    for _ in range(60): I = P/V; V = V0 - I*R
    a = (rpm/(KV*V) - 0.15)/0.8
    return (tho(a) if a <= 1 else None), I, P, V, KAP_AH*KULL/I*60
bas = ["Pil V_açık", "AUW kg", "eklenen yük kg", "askı gazı ThO", "askı akımı A", "askı gücü W", "V_yük V", "%80 pil ile dk", "hüküm"]
satirlar = []
for V0, ad in ((16.8, "dolu 4,20 V/hücre"), (15.6, "F logu başı 3,90 V/hücre")):
    for m in (4.0, 4.5, 5.0, 5.5, 6.0, 6.5, 7.0, 7.5, 8.0):
        t, I, P, V, dk = aski(m, V0)
        h = ("kalkamaz" if t is None else "tehlikeli (>0,70)" if t > 0.70 else "sınırda (0,55–0,70)" if t > 0.55 else "uygun")
        satirlar.append([f"{V0:.1f} ({ad})", m, round(m-AUW0, 1), (round(t, 2) if t is not None else "—"), round(I, 1), round(P), round(V, 2), (round(dk, 1) if t is not None else "—"), h])
NOT = ("Kaynak: performans_zarfi/tablolar.json kalibrasyonu (askı 0,36 @ 14,0 V / 37,3 A; tam gaz T/W 1,91 @ 12,36 V / 115 A); AUW 4,0 kg kullanıcı beyanı (15 Eyl); "
       "R_paket 30,4 mΩ; kapasite 13 Ah × %80. Ölçülen azami T/W 1,91 → azami AUW 7,6 kg (sarkmış pil), dolu pil modeli 8,5 kg. Hüküm eşikleri: ArduPilot askı gazı >0,70 yetersiz pay.")
X = os.path.join(os.path.dirname(os.path.abspath(__file__)), "NAVITAS_ucus_analizi_tablolar.xlsx")
wb = openpyxl.load_workbook(X)
if "Y_Yuk_Aski_Gazi" in wb.sheetnames: del wb["Y_Yuk_Aski_Gazi"]
ws = wb.create_sheet("Y_Yuk_Aski_Gazi")
ws.append(["Y — 4S 13 Ah + 9x5x3: eklenen yük → askı gazı / akım / uçuş süresi (AUW 4,0 kg)"]); ws["A1"].font = Font(bold=True, size=12)
ws.append(bas)
for c in ws[2]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDE4F0"); c.alignment = Alignment(wrap_text=True)
for r in satirlar: ws.append(r)
ws.append([]); ws.append([NOT]); ws.freeze_panes = "A3"
for col, w in zip("ABCDEFGHI", (24, 8, 14, 13, 12, 11, 9, 14, 20)): ws.column_dimensions[col].width = w
idx = wb["INDEKS"]
if not any((row[0].value == "Y_Yuk_Aski_Gazi") for row in idx.iter_rows(min_row=2)):
    idx.append(["Y_Yuk_Aski_Gazi", "Eklenen yük → askı gazı/akım/uçuş süresi (4S 13 Ah + 9x5x3, AUW 4,0 kg)", "performans_zarfi/tablolar.json kalibrasyonu + yuk_aski_tablosu.py", len(satirlar), len(bas), "46_Y_Yuk_Aski_Gazi.png"])
    idx.cell(row=idx.max_row, column=1).hyperlink = "#'Y_Yuk_Aski_Gazi'!A1"
wb.save(X)
fig, ax = plt.subplots(figsize=(15, 8.2)); ax.axis("off")
tb = ax.table(cellText=[[str(x).replace(".", ",") if isinstance(x, float) else str(x) for x in r] for r in satirlar], colLabels=bas, loc="upper center", cellLoc="center")
tb.auto_set_font_size(False); tb.set_fontsize(9); tb.scale(1, 1.35)
for (i, j), c in tb.get_celld().items():
    if i == 0: c.set_text_props(weight="bold"); c.set_facecolor("#DDE4F0")
    elif j == 8:
        c.set_facecolor({"uygun": "#DDF5DD", "sınırda (0,55–0,70)": "#FFF2CC", "tehlikeli (>0,70)": "#F8D0D0", "kalkamaz": "#E0E0E0"}[satirlar[i-1][8]])
ax.set_title("4S 13 Ah + 9x5x3 — eklenen yük → askı gazı, akım, uçuş süresi (AUW 4,0 kg)", loc="left", fontsize=12, weight="bold")
fig.text(0.01, 0.01, NOT, fontsize=7, wrap=True)
fig.savefig(os.path.join(os.path.dirname(os.path.abspath(__file__)), "png", "46_Y_Yuk_Aski_Gazi.png"), dpi=170, bbox_inches="tight"); print("ok", len(satirlar), "satir")
